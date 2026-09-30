from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
E = ROOT / 'docs' / 'evidence'
doc = Document()
doc.styles['Normal'].font.size = Pt(10.5)


def table(head, rows):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(head):
        t.rows[0].cells[i].text = h
    for r in rows:
        c = t.add_row().cells
        for i, v in enumerate(r):
            c[i].text = v


def fig(name, cap):
    p = E / name
    if not p.exists():
        doc.add_paragraph(f'SCREENSHOT NOT AVAILABLE: {name}').runs[0].bold = True
        return
    doc.add_picture(str(p), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(cap).runs[0].italic = True


doc.add_heading('IT2406 Performance Task 1: PS and FI, Development of Ultralight Bike', 0)
doc.add_paragraph('SAP system M53, client 236, host m53p.ucc.cloud. SAP user LEARN-653. Work dates 2026-09-29 and 2026-09-30.')
doc.add_paragraph('Every task was run by the Ultralight control site (ultralight-project-builder.pages.dev) and its Playwright runner, in Autopilot mode. Each result below was read back from SAP.')

doc.add_heading('Key reference numbers', 1)
table(['Item', 'Value'], [
    ('Project definition', 'P/2653, Development of Ultralight Bike 653'),
    ('WBS elements', 'P/2653, P/2653-1 to P/2653-5, controlling area EU00'),
    ('Network', '4000100, activities 0010 to 0140 plus 0045 (external) and 0135 (costs)'),
    ('PS text', 'PH-653-1'),
    ('Milestones', '255 (activity 0070), 256 (0120), 257 (0140)'),
    ('Project status', 'REL (released)'),
    ('Supplier', '114653 Pyramid Biking'),
    ('Supplier invoice (FB60)', 'Document 1900000063, 9,700.00 EUR, G/L 6300000, network 4000100 activity 0135'),
    ('Confirmation (CN25)', 'Activity 0010, 35 h actual of 80 h'),
])

doc.add_heading('Status per task', 1)
table(['Task', 'Status', 'How it was checked'], [
    ('1 Project, WBS, responsible cost centres', 'Done', 'CJ20N Responsibilities tab, Figure 1'),
    ('2 Activities, 0045 external services, 0135 primary costs', 'Done', 'CJ20N Activity Overview, Figure 2'),
    ('3 Network graph before relationships', 'Done, see note', 'Figure 3'),
    ('4 22 finish-to-start relationships', 'Done', 'Network graph shows the links, Figure 4'),
    ('5 Network graph after relationships', 'Done', 'Figure 4'),
    ('6 PS text and milestones', 'Done', 'Project tree shows PH-653-1 and milestones 255, 256'),
    ('7 Release project', 'Done', 'System Status REL read in CJ20N (Validate run #11)'),
    ('8 Planned cost report', 'Done', 'S_ALR_87013542, Figure 5'),
    ('9 Structure overview', 'Done', 'CN41N, Figure 6'),
    ('10 Primary cost 0135 to 8,000 EUR, flexible', 'Done', 'Plan on 6300000 changed from 10,000.00 to 8,000.00 (Figures 5 and 7)'),
    ('11 Confirm 35 h on 0010', 'Done', 'Labor actual 1,750.00 on 8000000, Figure 7'),
    ('12 Cost report after confirmation', 'Done', 'Figure 7'),
    ('13 Supplier invoice 9,700 EUR', 'Done', 'FB60 entry, Appendix A1; actual 9,700.00 on 6300000, Figure 8'),
    ('14 Final cost report', 'Done', 'Figure 8'),
])

doc.add_heading('Screenshots', 1)
fig('t1-wbs-responsibilities.jpg', 'Figure 1. Task 1, CJ20N, Responsibilities tab: WBS P/2653 and P/2653-1 to -5, controlling area EU00, responsible cost centres EURD1000, EURD1000, EURD1000, EUQM1000, EUPR1000, EUQM1000.')
fig('t2-activity-overview.jpg', 'Figure 2. Task 2, CJ20N, Activity Overview: activities 0010 to 0140 assigned to their WBS elements, status CRTD before release.')
fig('job9-t3-network-before.jpg', 'Figure 3. Task 3, Fiori app Project Network Graph. Note: this capture comes from Autopilot run #9, which ran after the relationships already existed, so the links are visible. No capture of the true "before" state was saved when Task 3 first ran.')
fig('job9-t5-network-after.jpg', 'Figure 4. Tasks 4 and 5, Project Network Graph after the relationships: 0010 leads to 0020, 0030, 0040 and 0045; 0020 to 0050; 0030 to 0060; 0040, 0045, 0050 and 0060 to 0070 Prototype assembly.')
fig('t8-costs-planned.jpg', 'Figure 5. Task 8, S_ALR_87013542 on 2026-09-29, after release and before Task 10: no actual costs, commitment 5,000.00, plan 51,433.14 EUR (6300000 plan 10,000.00).')
fig('job18-t9-structure.jpg', 'Figure 6. Task 9, CN41N Structure Overview with PS info profile GL01000: P/2653, PH-653-1, network 4000100 and WBS P/2653-1 to -5. Captured by Autopilot run #18 on 2026-09-30.')
fig('t12-costs-after-confirmation.jpg', 'Figure 7. Task 12, S_ALR_87013542 after Tasks 10 and 11: actual 1,750.00 EUR on 8000000 Labor (35 h), plan 49,433.14 EUR (6300000 plan now 8,000.00).')
fig('t14-costs-final.jpg', 'Figure 8. Task 14, S_ALR_87013542 after the supplier invoice: actual 9,700.00 on 6300000 plus 1,750.00 on 8000000, total actual 11,450.00 EUR, commitment 5,000.00, total 16,450.00, plan 49,433.14.')

doc.add_heading('Appendix', 1)
fig('t13-before-post.jpg', 'Appendix A1. Task 13, FB60 before posting: supplier 114653, amount 9,700.00 EUR, G/L 6300000, network 4000100, activity 0135, balance 0.00. SAP posted it as document 1900000063.')
fig('job9-t2-activity-overview.jpg', 'Appendix A2. CJ20N project tree after all tasks: PS text PH-653-1, network 4000100 with 0045 and 0135, milestones 255 and 256.')

doc.add_heading('Conclusion', 1)
table(['Cost element', 'Plan, Task 8', 'Plan, Task 12', 'Actual, Task 12', 'Actual, Task 14'], [
    ('6300000 Other operating expenses', '10,000.00', '8,000.00', '0.00', '9,700.00'),
    ('6991000 Cost of Labor (commitment 5,000.00)', '5,000.00', '5,000.00', '0.00', '0.00'),
    ('8000000 Labor', '21,250.00', '21,250.00', '1,750.00', '1,750.00'),
    ('8012000, 8013000, 8014000 Overhead', '15,183.14', '15,183.14', '0.00', '0.00'),
    ('All cost elements', '51,433.14', '49,433.14', '1,750.00', '11,450.00'),
])
doc.add_paragraph('')
doc.add_paragraph('The two planned reports differ only on 6300000. Task 10 cut the primary cost on activity 0135 from 10,000 to 8,000 EUR, so the total plan fell by exactly 2,000 EUR, from 51,433.14 to 49,433.14. Labor and overhead plans did not change.')
doc.add_paragraph('The first actual cost was the Task 11 confirmation: 35 hours on activity 0010 gave 1,750.00 EUR on 8000000 Labor, which is 50 EUR per hour. The Task 13 supplier invoice then added 9,700.00 EUR on 6300000, for 11,450.00 EUR total actual.')
doc.add_paragraph('On 6300000 the actual 9,700.00 is 1,700.00 EUR above the new 8,000.00 plan. It is still 300.00 EUR below the original 10,000.00 plan, so the Task 10 cut was too deep for this activity. The project as a whole has used 11,450.00 of 49,433.14 EUR planned (about 23 %). The 5,000.00 EUR commitment for the outsourced gear (0045) is still open. No overhead has been charged yet, because overhead is applied at period-end costing and that is not part of this task.')

doc.add_heading('Not confirmed', 1)
for s in [
    'Grading monitor percentage for LEARN-653. The grading app is not assigned to this user, so only the student can read it. Where I looked: the task notes (pasted-text, section on the grading monitor).',
    'A capture of the network graph with no links (Task 3 "before"). The only saved capture shows the links (see Figure 3).',
    'That milestone 257 on activity 0140 appears in the tree. It was read back during the run, but it is below the visible area in Appendix A2.',
]:
    doc.add_paragraph(s, style='List Bullet')

out = ROOT / 'IT2406_Performance_Task_1_Ultralight_Bike_LEARN-653.docx'
doc.save(out)
print(out)
